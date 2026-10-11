from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("TenantAI_Jury_Presentation.pptx")
WIDTH = 13.333
HEIGHT = 7.5
BG = RGBColor(7, 13, 27)
PANEL = RGBColor(16, 27, 48)
PANEL_2 = RGBColor(20, 36, 62)
TEXT = RGBColor(239, 246, 255)
MUTED = RGBColor(156, 173, 197)
TEAL = RGBColor(45, 212, 191)
BLUE = RGBColor(96, 165, 250)
AMBER = RGBColor(251, 191, 36)
RED = RGBColor(248, 113, 113)
LINE = RGBColor(46, 67, 98)


def text(slide, value, x, y, w, h, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def panel(slide, x, y, w, h, title, body, accent=TEAL, size=12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    text(slide, title, x + 0.2, y + 0.17, w - 0.4, 0.3, 14, accent, True)
    text(slide, body, x + 0.2, y + 0.62, w - 0.4, h - 0.76, size, TEXT)


def metric(slide, x, y, w, value, label, accent=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.1)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL_2
    shape.line.color.rgb = LINE
    text(slide, value, x + 0.12, y + 0.12, w - 0.24, 0.4, 24, accent, True, PP_ALIGN.CENTER)
    text(slide, label, x + 0.12, y + 0.62, w - 0.24, 0.25, 9, MUTED, False, PP_ALIGN.CENTER)


def base(prs, title, subtitle="", section="PRODUCT"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(HEIGHT)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = TEAL
    text(slide, section, 0.72, 0.25, 2, 0.22, 9, TEAL, True)
    text(slide, title, 0.68, 0.52, 12, 0.52, 27, TEXT, True)
    if subtitle:
        text(slide, subtitle, 0.7, 1.08, 11.8, 0.3, 11, MUTED)
    number = len(prs.slides)
    text(slide, f"tenantAI  |  TriCity AI Hackathon  |  {number:02d}", 0.7, 7.16, 12, 0.18, 8, MUTED)
    return slide


def add_flow(slide, labels, y=2.4, start=0.65, width=1.45, gap=0.18):
    for index, label in enumerate(labels):
        x = start + index * (width + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(0.9)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL_2
        shape.line.color.rgb = TEAL
        text(slide, str(index + 1), x + 0.08, y + 0.1, 0.22, 0.2, 10, TEAL, True)
        text(slide, label, x + 0.28, y + 0.18, width - 0.36, 0.48, 11, TEXT, True, PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = Inches(WIDTH)
prs.slide_height = Inches(HEIGHT)

# 01 - Cover
slide = base(prs, "tenantAI", "Secure multi-tenant AI knowledge platform", "PITCH")
text(slide, "Company knowledge.\nGrounded answers.\nProvable isolation.", 0.9, 1.85, 7.2, 2.2, 31, TEXT, True)
text(slide, "An enterprise assistant that turns private documents into cited answers while enforcing the workspace boundary in the backend.", 0.95, 4.35, 7.2, 0.8, 17, MUTED)
metric(slide, 9.0, 2.0, 1.65, "RAG", "grounded answers", TEAL)
metric(slide, 10.85, 2.0, 1.65, "RBAC", "role controls", BLUE)
metric(slide, 9.0, 3.35, 1.65, "404", "hidden IDOR", AMBER)
metric(slide, 10.85, 3.35, 1.65, "OTP", "verified identity", TEAL)
text(slide, "TriCity AI Hackathon | SaaSsum AI track | jury-ready MVP", 0.95, 6.3, 8.5, 0.3, 12, TEAL, True)

# 02 - One minute pitch
slide = base(prs, "The one-minute pitch", "What tenantAI does and why it matters", "PITCH")
panel(slide, 0.75, 1.7, 3.8, 3.9, "THE PAIN", "Company knowledge is trapped in PDFs, manuals, and policies. Teams lose time searching, answers are hard to verify, and naive shared AI search can leak one company's information to another.", RED, 14)
panel(slide, 4.8, 1.7, 3.8, 3.9, "THE PRODUCT", "tenantAI gives each company a private knowledge workspace. Users upload approved material, ask natural-language questions, receive source citations, and see the answer path through audit telemetry.", TEAL, 14)
panel(slide, 8.85, 1.7, 3.8, 3.9, "THE PROOF", "The jury can switch between TechNova and GreenLeaf, compare answers, attempt a cross-tenant resource read, observe HTTP 404 with zero leaked bytes, and inspect the security event.", AMBER, 14)

# 03 - Problem
slide = base(prs, "Why current enterprise search breaks", "The problem combines productivity, trust, and security", "PROBLEM")
panel(slide, 0.75, 1.75, 2.85, 3.8, "01  FINDABILITY", "Important knowledge is distributed across manuals, policies, and operational files.\n\nPeople ask the same questions repeatedly.", BLUE)
panel(slide, 3.75, 1.75, 2.85, 3.8, "02  CONFIDENCE", "A fluent answer without a source is difficult to approve, explain, or correct.\n\nThe cost is slow decisions.", AMBER)
panel(slide, 6.75, 1.75, 2.85, 3.8, "03  PRIVACY", "A shared vector search without strict ownership filters can return another company's content.\n\nOne missing predicate is enough.", RED)
panel(slide, 9.75, 1.75, 2.85, 3.8, "04  GOVERNANCE", "Without audit events, operators cannot explain who asked what, which tenant was used, or whether a security probe was blocked.", TEAL)
text(slide, "tenantAI treats tenant boundaries as a core product capability, not a frontend checkbox.", 1.0, 6.15, 11.3, 0.35, 15, TEAL, True, PP_ALIGN.CENTER)

# 04 - Solution
slide = base(prs, "The solution in one sentence", "A verified, tenant-scoped, citation-first AI workspace", "PRODUCT")
text(slide, "Every answer follows the same contract:", 0.95, 1.75, 6, 0.4, 19, TEXT, True)
add_flow(slide, ["Verify identity", "Resolve tenant", "Retrieve owned chunks", "Ground answer", "Cite or refuse", "Write audit"])
panel(slide, 1.0, 4.15, 3.45, 1.4, "USEFUL", "Natural-language search over company documents.", TEAL, 14)
panel(slide, 4.95, 4.15, 3.45, 1.4, "EXPLAINABLE", "Answers carry document/page evidence.", BLUE, 14)
panel(slide, 8.9, 4.15, 3.45, 1.4, "DEFENSIBLE", "Ownership and denial are testable server behavior.", AMBER, 14)

# 05 - Users
slide = base(prs, "Who uses tenantAI?", "A shared platform with separate responsibilities", "PRODUCT")
panel(slide, 0.75, 1.8, 3.8, 3.6, "TENANT ADMINISTRATOR", "Creates the workspace, manages members, reviews quotas, and inspects audit telemetry.\n\nJury value: shows governance, not just a chatbot.", TEAL, 14)
panel(slide, 4.8, 1.8, 3.8, 3.6, "KNOWLEDGE WORKER", "Uploads approved documents, asks questions, follows citations, and revisits private conversations.\n\nJury value: shows daily user utility.", BLUE, 14)
panel(slide, 8.85, 1.8, 3.8, 3.6, "PLATFORM SUPERADMIN", "Reviews global tenant usage, provisions workspaces, and runs the controlled IDOR security simulation.\n\nJury value: shows SaaS operations.", AMBER, 14)

# 06 - Product surfaces
slide = base(prs, "Product surfaces", "The MVP covers the complete knowledge-work loop", "PRODUCT")
panel(slide, 0.65, 1.7, 2.85, 4.1, "IDENTITY", "Registration\nEmail OTP\nVerified login\nJWT session\nRole-based access\n\nUnverified accounts cannot enter protected APIs.", TEAL, 12)
panel(slide, 3.65, 1.7, 2.85, 4.1, "KNOWLEDGE", "PDF/TXT/MD uploads\nValidation and quotas\nText extraction\nChunking and embeddings\nPENDING / READY / FAILED\n\nOnly READY content is searchable.", BLUE, 12)
panel(slide, 6.65, 1.7, 2.85, 4.1, "ASSISTANT", "Tenant-scoped RAG\nCitations\nConversation history\nOut-of-domain refusal\nConfidence and latency\n\nUseful answers remain traceable.", TEAL, 12)
panel(slide, 9.65, 1.7, 2.85, 4.1, "CONTROL", "Audit logs\nUsage telemetry\nTenant quotas\nAdmin center\nIDOR simulator\n\nSecurity is visible to operators.", AMBER, 12)

# 07 - Architecture
slide = base(prs, "Architecture: one system, one source of truth", "The browser never decides which tenant's data it may read", "ENGINEERING")
panel(slide, 0.7, 1.75, 2.6, 1.45, "NEXT.JS", "Accessible dashboard, auth, chat, documents, admin views", BLUE)
panel(slide, 3.65, 1.75, 2.6, 1.45, "FASTAPI", "JWT verification, RBAC, tenant guards, API contracts", TEAL)
panel(slide, 6.6, 1.75, 2.6, 1.45, "SQLALCHEMY", "Users, tenants, documents, chunks, conversations, logs", AMBER)
panel(slide, 9.55, 1.75, 2.6, 1.45, "RAG ENGINE", "Scoped retrieval, grounded synthesis, citations, refusal", TEAL)
add_flow(slide, ["Browser", "Auth gateway", "Tenant filter", "Private data", "RAG", "Audited response"], y=4.05, start=0.55, width=1.72, gap=0.18)
text(slide, "Invariant: tenant_id comes from verified server-side identity. Client-supplied tenant IDs are ignored for regular users.", 1.0, 5.75, 11.2, 0.55, 16, TEXT, True, PP_ALIGN.CENTER)

# 08 - Data model
slide = base(prs, "Data model: where isolation lives", "Every private relationship has an ownership boundary", "ENGINEERING")
panel(slide, 0.5, 1.75, 2.35, 4.0, "TENANT", "Company workspace\nPlan and quotas\nDomain\nOwns private records", TEAL, 12)
panel(slide, 3.0, 1.75, 2.35, 4.0, "USER", "Verified email\nPassword hash\nRole\nBelongs to tenant", BLUE, 12)
panel(slide, 5.5, 1.75, 2.35, 4.0, "DOCUMENT", "Tenant-owned file\nParser status\nSource metadata\nSafe deletion scope", AMBER, 12)
panel(slide, 8.0, 1.75, 2.35, 4.0, "CHUNK", "Text segment\nEmbedding vector\nPage metadata\nTenant-scoped search", TEAL, 12)
panel(slide, 10.5, 1.75, 2.35, 4.0, "CONVERSATION", "Owner user\nTenant relationship\nMessages and citations\nPrivate history", BLUE, 12)
text(slide, "A valid numeric ID is never enough: the authenticated tenant and, for conversations, the owning user must also match.", 0.8, 6.15, 11.8, 0.4, 14, TEAL, True, PP_ALIGN.CENTER)

# 09 - RAG
slide = base(prs, "How grounded RAG works", "The assistant is constrained by approved workspace evidence", "AI")
add_flow(slide, ["Question", "Embed", "Tenant filter", "Rank chunks", "Ground prompt", "Citations"])
panel(slide, 0.9, 4.05, 3.55, 1.55, "IN-DOMAIN", "TechNova: “What is the battery warranty?”\n\nReturns an answer with the hardware manual citation.", TEAL, 13)
panel(slide, 4.9, 4.05, 3.55, 1.55, "OUT-OF-DOMAIN", "GreenLeaf: “What is the UltraBook warranty?”\n\nReturns a refusal with no unauthorized citations.", AMBER, 13)
panel(slide, 8.9, 4.05, 3.55, 1.55, "TRACEABILITY", "The response preserves source title, page, excerpt, confidence, latency, and token usage for review.", BLUE, 13)

# 10 - Security layers
slide = base(prs, "Defense in depth", "No single control carries the whole security model", "SECURITY")
panel(slide, 0.75, 1.7, 2.85, 3.95, "IDENTITY", "Email OTP verification\nBcrypt password hashing\nJWT algorithm allowlist\nActive-user check\n\nStops anonymous and unverified access.", TEAL, 12)
panel(slide, 3.75, 1.7, 2.85, 3.95, "AUTHORIZATION", "SuperAdmin and tenant-admin guards\nOwner checks for conversations\nTenant predicates on resources\nHidden 404 responses\n\nStops horizontal and vertical access.", BLUE, 12)
panel(slide, 6.75, 1.7, 2.85, 3.95, "DATA", "READY-only retrieval\nPrivate storage direction\nSafe filename handling\nGeneric processing errors\n\nReduces unsafe content and leakage.", AMBER, 12)
panel(slide, 9.75, 1.7, 2.85, 3.95, "OBSERVABILITY", "Audit events\nLatency/tokens\nSecurity simulation\nRedacted admin logs\nSecurity response headers\n\nMakes controls provable.", TEAL, 12)

# 11 - IDOR proof
slide = base(prs, "The security moment: prove the IDOR block", "A jury can observe the control in real time", "SECURITY")
text(slide, "ATTACKER", 0.85, 1.75, 2.5, 0.3, 15, RED, True)
text(slide, "GreenLeaf attempts to fetch TechNova document #1.", 0.85, 2.15, 3.2, 0.9, 20, TEXT, True)
text(slide, "ENFORCEMENT", 4.95, 1.75, 2.5, 0.3, 15, BLUE, True)
text(slide, "JWT tenant identity overrides client input. Query requires matching tenant ownership.", 4.95, 2.15, 3.3, 1.2, 18, TEXT, True)
text(slide, "RESULT", 9.05, 1.75, 2.5, 0.3, 15, TEAL, True)
text(slide, "HTTP 404\n0 bytes leaked\nAudit event written", 9.05, 2.15, 3.2, 1.3, 23, TEXT, True)
panel(slide, 1.0, 4.65, 11.2, 1.0, "WHY 404?", "A hidden 404 avoids confirming that another tenant's resource exists. The denial is also recorded for the tenant's security trail.", AMBER, 14)

# 12 - Auth lifecycle
slide = base(prs, "Authentication lifecycle", "Trust is established before the first protected request", "SECURITY")
add_flow(slide, ["Register", "Generate OTP", "Hash + expire", "Verify email", "Issue JWT", "Check active user"])
panel(slide, 1.0, 4.25, 3.45, 1.45, "OTP CONTROLS", "Cryptographically random code, hashed storage, five-minute expiry, attempt limit, resend cooldown.", TEAL, 12)
panel(slide, 4.95, 4.25, 3.45, 1.45, "SESSION CONTROLS", "JWT algorithm allowlist, active-user lookup, production secret length validation.", BLUE, 12)
panel(slide, 8.9, 4.25, 3.45, 1.45, "USER EXPERIENCE", "Clear error messages, development-log OTP fallback, and no silent success on email failure.", AMBER, 12)

# 13 - Upload pipeline
slide = base(prs, "Document pipeline", "From file upload to searchable evidence", "ENGINEERING")
add_flow(slide, ["Upload", "Validate", "Extract", "Chunk", "Embed", "Index", "Retrieve", "Cite"], y=2.15, start=0.48, width=1.45, gap=0.16)
text(slide, "Supported demo formats: PDF, TXT, and MD", 0.95, 4.2, 5.4, 0.4, 17, TEAL, True)
text(slide, "Failure behavior is explicit: unsupported type, size limit, quota, parser failure, or embedding failure.", 0.95, 4.8, 7.6, 0.55, 15, TEXT)
panel(slide, 9.1, 4.05, 3.35, 1.55, "STATUS MODEL", "PENDING → PROCESSING → READY\nor FAILED with safe user-facing messaging.", AMBER, 13)

# 14 - Telegram
slide = base(prs, "One knowledge layer, multiple channels", "The Telegram worker reuses the same tenant-aware behavior", "INTEGRATION")
panel(slide, 0.75, 1.75, 3.8, 3.85, "AUTHENTICATED SESSION", "/start and login establish a tenant-aware session.\n\nThe inactivity window is sliding: authenticated activity refreshes it; 120 seconds without activity expires it.", TEAL, 14)
panel(slide, 4.8, 1.75, 3.8, 3.85, "KNOWLEDGE COMMANDS", "/chat\n/documents\n/upload\n/conversations\n/audits\n/telemetry\n\nThe bot exposes the same document and chat workflow.", BLUE, 14)
panel(slide, 8.85, 1.75, 3.8, 3.85, "ADMIN COMMANDS", "/settings\n/setcompany\n/setdomain\n\nAdmin-only changes, explicit upload results, generic processing errors, and command menu registration.", AMBER, 13)

# 15 - Observability
slide = base(prs, "Observability: security becomes evidence", "Operators can answer what happened, where, and why", "GOVERNANCE")
metric(slide, 0.9, 1.8, 2.25, "WHO", "user / role", TEAL)
metric(slide, 3.45, 1.8, 2.25, "WHERE", "tenant", BLUE)
metric(slide, 6.0, 1.8, 2.25, "WHAT", "action / outcome", AMBER)
metric(slide, 8.55, 1.8, 2.25, "HOW FAST", "latency", TEAL)
metric(slide, 11.1, 1.8, 1.3, "COST", "tokens", BLUE)
panel(slide, 1.0, 3.65, 5.35, 1.75, "SECURITY EVENTS", "Login, upload, deletion, chat, denied resource access, and IDOR simulation are recorded. Prompt content is not included in chat telemetry.", TEAL, 14)
panel(slide, 6.95, 3.65, 5.35, 1.75, "PRIVACY BOUNDARY", "Audit access is limited to tenant administrators and SuperAdmins. Email, IP addresses, and chat prompt excerpts are redacted from responses.", BLUE, 14)

# 16 - Demo script
slide = base(prs, "The 7-minute jury demo", "A rehearsable sequence with a visible payoff", "DEMO")
steps = [
    ("01", "Open landing page", "Explain secure workspace promise."),
    ("02", "Login as TechNova", "Show private document list."),
    ("03", "Ask warranty question", "Show grounded response and citation."),
    ("04", "Switch to GreenLeaf", "Show a separate knowledge boundary."),
    ("05", "Ask TechNova question", "Show refusal with no foreign citation."),
    ("06", "Run IDOR simulator", "Show 404 and zero leaked bytes."),
    ("07", "Open telemetry", "Show the recorded security event."),
]
for i, (num, title, body) in enumerate(steps):
    y = 1.55 + i * 0.7
    text(slide, num, 0.9, y, 0.5, 0.3, 14, TEAL, True)
    text(slide, title, 1.55, y, 3.2, 0.3, 15, TEXT, True)
    text(slide, body, 5.0, y, 7.0, 0.3, 13, MUTED)

# 17 - Credentials
slide = base(prs, "Jury demo fixtures", "Use only for demonstration; rotate before real users or production", "DEMO")
panel(slide, 0.75, 1.8, 3.8, 3.8, "TECHNOVA SOLUTIONS", "Admin: admin@technova.com\nPassword: DemoPass123!\n\nKnowledge: UltraBook Pro X15 hardware manual\n\nAsk: “What is the battery warranty?”", TEAL, 14)
panel(slide, 4.8, 1.8, 3.8, 3.8, "GREENLEAF RETAIL", "Admin: admin@greenleaf.com\nPassword: DemoPass123!\n\nKnowledge: customer refund and return policy\n\nAsk: “What is the return window?”", BLUE, 14)
panel(slide, 8.85, 1.8, 3.8, 3.8, "PRESENTATION RULE", "Do not use these credentials in production.\n\nUse them to demonstrate isolation, citations, refusal behavior, and audit evidence.\n\nThe requested fixtures remain unchanged.", AMBER, 14)

# 18 - Testing
slide = base(prs, "Verification before the pitch", "The product is tested at the same boundaries the jury will inspect", "QUALITY")
metric(slide, 0.95, 1.75, 2.45, "23", "backend tests passing", TEAL)
metric(slide, 3.75, 1.75, 2.45, "14", "frontend routes built", BLUE)
metric(slide, 6.55, 1.75, 2.45, "0", "bounds issues in deck", AMBER)
metric(slide, 9.35, 1.75, 2.45, "404", "hidden IDOR result", TEAL)
panel(slide, 1.0, 3.65, 3.55, 1.7, "REGRESSION TESTS", "Auth verification, tenant document isolation, RAG citations, admin access, conversation privacy, and Telegram session expiry.", TEAL, 13)
panel(slide, 4.9, 3.65, 3.55, 1.7, "FRONTEND CHECKS", "Production TypeScript build, accessible controls, responsive flows, reduced-motion support, and role-aware navigation.", BLUE, 13)
panel(slide, 8.8, 3.65, 3.55, 1.7, "HONEST STATUS", "Deployment verification, object storage, OCR, async processing, and formal migrations remain roadmap work.", AMBER, 13)

# 19 - Deployment
slide = base(prs, "From local proof to Render", "Production is a controlled release process", "OPERATIONS")
add_flow(slide, ["Create blueprint", "Provision DB", "Set secrets", "Deploy API", "Deploy UI", "Smoke test"])
panel(slide, 0.9, 4.15, 3.6, 1.55, "RENDER SERVICES", "Frontend, FastAPI backend, PostgreSQL, and optional Telegram worker.", TEAL, 13)
panel(slide, 4.85, 4.15, 3.6, 1.55, "REQUIRED SECRETS", "SECRET_KEY, DATABASE_URL, CORS_ORIGINS, SMTP, AI provider keys, Telegram token.", BLUE, 13)
panel(slide, 8.8, 4.15, 3.6, 1.55, "SMOKE TEST", "Health → register → verify → login → upload → chat → logs → isolation probe.", AMBER, 13)

# 20 - Security hardening
slide = base(prs, "Security hardening shipped", "The implementation closes common multi-tenant failure modes", "SECURITY")
panel(slide, 0.75, 1.7, 2.85, 4.0, "SECRETS", "Removed hardcoded database and JWT secrets from Docker Compose.\n\nProduction now requires environment-managed values and a strong signing key.", TEAL, 12)
panel(slide, 3.75, 1.7, 2.85, 4.0, "OWNERSHIP", "Conversation read, append, and delete operations require both tenant ownership and owning user.\n\nSame-tenant users cannot hijack threads.", BLUE, 12)
panel(slide, 6.75, 1.7, 2.85, 4.0, "EXPOSURE", "Logs are admin-only and redacted. Production docs are disabled. Health errors avoid raw internals. Security headers are enabled.", AMBER, 12)
panel(slide, 9.75, 1.7, 2.85, 4.0, "FRONTEND", "Navigation hides audit/admin surfaces from unauthorized roles and avoids unnecessary log requests for regular users.", TEAL, 12)

# 21 - Business
slide = base(prs, "Why this can become a SaaS product", "Security and utility create a credible expansion path", "BUSINESS")
panel(slide, 0.75, 1.75, 3.8, 3.8, "CUSTOMER VALUE", "Reduce repeated support questions.\nShorten onboarding.\nFind policy answers faster.\nKeep evidence attached to decisions.\n\nThe buyer gets productivity without giving up tenant boundaries.", TEAL, 14)
panel(slide, 4.8, 1.75, 3.8, 3.8, "PRICING SURFACES", "Per workspace\nPer seat\nStorage and document volume\nAI query usage\nPremium audit and retention\n\nUsage telemetry makes these surfaces measurable.", BLUE, 14)
panel(slide, 8.85, 1.75, 3.8, 3.8, "EXPANSION", "OCR for scanned files\nAsync workers\nPrivate object storage\nSlack/Teams/Telegram channels\nSAML/SCIM and enterprise policy controls", AMBER, 14)

# 22 - Roadmap
slide = base(prs, "Roadmap: from MVP to production platform", "Prioritized by reliability, scale, and enterprise readiness", "ROADMAP")
panel(slide, 0.75, 1.75, 3.8, 3.8, "NOW", "Working tenant-scoped APIs\nGrounded RAG and citations\nOTP auth and RBAC\nAudit telemetry\nIDOR simulator\nAccessible web UI", TEAL, 14)
panel(slide, 4.8, 1.75, 3.8, 3.8, "NEXT", "Alembic migrations\nPrivate object storage\nRate limiting and WAF\nDependency scanning\nVerified Render rollout\nOperational alerting", BLUE, 14)
panel(slide, 8.85, 1.75, 3.8, 3.8, "LATER", "OCR and async ingestion\nSAML/SCIM\nRetention policies\nFeedback and evaluation sets\nTenant-level model policies\nEnterprise integrations", AMBER, 14)

# 23 - Q&A
slide = base(prs, "Jury questions: concise answers", "Use these as speaking prompts", "CLOSE")
panel(slide, 0.65, 1.65, 2.85, 4.15, "WHY MULTI-TENANT?", "Because a company assistant must be useful without mixing confidential workspaces. Tenant scope is enforced in backend queries.", TEAL, 13)
panel(slide, 3.65, 1.65, 2.85, 4.15, "WHY RAG?", "Company knowledge changes. Retrieval grounds answers in approved documents and provides evidence instead of relying on model memory.", BLUE, 13)
panel(slide, 6.65, 1.65, 2.85, 4.15, "HOW PROVE SECURITY?", "Run the IDOR simulator, observe 404 and zero leaked bytes, then inspect the audit event and tenant boundary.", AMBER, 13)
panel(slide, 9.65, 1.65, 2.85, 4.15, "WHAT NEXT?", "Complete production infrastructure: migrations, storage, async processing, OCR, rate limiting, and verified deployment.", TEAL, 13)

# 24 - Close
slide = base(prs, "The takeaway", "Make the answer useful. Make the boundary provable.", "CLOSE")
text(slide, "tenantAI is not only a chatbot.", 1.0, 1.8, 11.3, 0.55, 28, TEXT, True, PP_ALIGN.CENTER)
text(slide, "It is a secure knowledge operating layer for companies that need speed, evidence, and separation.", 1.2, 2.7, 10.9, 0.8, 22, TEAL, True, PP_ALIGN.CENTER)
metric(slide, 2.0, 4.45, 2.3, "USEFUL", "grounded answers", TEAL)
metric(slide, 5.5, 4.45, 2.3, "SAFE", "tenant boundary", BLUE)
metric(slide, 9.0, 4.45, 2.3, "PROVABLE", "audit evidence", AMBER)
text(slide, "Thank you | Live demo begins with TechNova login", 1.0, 6.35, 11.3, 0.3, 14, MUTED, False, PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
